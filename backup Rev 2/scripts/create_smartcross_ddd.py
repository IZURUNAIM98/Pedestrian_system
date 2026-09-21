from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path("SmartCross_Detailed_Design_Description.docx")

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def borders(table, color="D9D9D9"):
    tblPr = table._tbl.tblPr
    edges = tblPr.first_child_found_in("w:tblBorders")
    if edges is None:
        edges = OxmlElement("w:tblBorders")
        tblPr.append(edges)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = edges.find(qn("w:" + edge))
        if el is None:
            el = OxmlElement("w:" + edge)
            edges.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), color)

def cell_text(cell, text, bold=False, color="000000", size=8.2):
    cell.text = ""
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(text)); run.bold = bold; run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def make_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; borders(t)
    for i, h in enumerate(headers):
        cell_text(t.rows[0].cells[i], h, True, "FFFFFF", 8); shade(t.rows[0].cells[i], "163D4A")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell_text(cells[i], value)
            if ri % 2: shade(cells[i], "F3F7F6")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def para(doc, text):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.08; p.add_run(text)

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3); p.add_run(text)

doc = Document()
sec = doc.sections[0]; sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
styles = doc.styles
styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(9)
for name, size, color in (("Title",25,"000000"),("Heading 1",16,"163D4A"),("Heading 2",12,"176B62")):
    s = styles[name]; s.font.name = "Aptos"; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor.from_string(color)
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT; footer.add_run("SmartCross Detailed Design Description | Simulation-only controlled prototype").font.size = Pt(7)

doc.add_paragraph("SmartCross Detailed Design Description", style="Title")
p = doc.add_paragraph(); r = p.add_run("As built design for the SmartCross pedestrian crossing simulator"); r.font.size = Pt(13); r.font.color.rgb = RGBColor.from_string("176B62")
para(doc, "Document status: As-built baseline | Revision 1.0 | Date: 18 September 2026")
para(doc, "Purpose: Define the implemented SmartCross architecture, state machines, interfaces, safety rules, evidence records, and deployment boundary so operators, reviewers, maintainers, and commissioning stakeholders can understand what the prototype does and what remains outside its authority.")
make_table(doc, ["Document control", "Value"], [("System", "SmartCross smart pedestrian-crossing simulator"),("Operating modes", "Normal Crossing and School Crossing"),("Primary users", "MPAJ demonstration operators, planners, safety reviewers, testers"),("Source of truth", "Current repository implementation and verified production bundle"),("Operational status", "Simulation-only; no live CCTV, roadside sensors, signal controller, enforcement, or emergency dispatch")])

doc.add_heading("1 System purpose and scope", 1)
para(doc, "SmartCross is a deterministic browser-based decision-support prototype for demonstrating protected pedestrian crossing logic in Malaysian urban conditions. It joins a visual crossing scene, scenario catalogue, traffic and pedestrian motion, CCTV and radar representations, signal states, event timeline, incident review, and PDF reporting. It demonstrates a traceable response but does not operate field equipment or make an enforcement conclusion.")
para(doc, "The as-built scope includes one no-violation baseline and 20 shared violation cases in each crossing mode, seven chronological timeline stages, 40 metre mirrored approach monitoring, incoming and outgoing CCTV representations, protected WALK timing, degraded camera fallback, and validated emergency-priority simulation.")

doc.add_heading("2 Architectural overview", 1)
make_table(doc, ["Layer", "Implemented responsibility", "Primary source components"], [("Presentation", "Authenticated dashboard, navigation, crossing scene, controls, timeline, history, incident review, PDF action", "components/dashboard.tsx; components/crossing-view.tsx; components/event-log.tsx"),("Simulation domain", "Scenario selection, deterministic motion frames, signal and pedestrian states, confidence and safety responses", "lib/simulation.ts; lib/motion.ts"),("Priority services", "Camera degradation fallback and emergency responder priority state machine", "lib/camera-fallback.ts; lib/emergency-priority.ts"),("Evidence and reporting", "Action-synchronised event log, incident records, synthetic plate evidence, PDF report generation", "lib/pdf-report.ts; components/event-log.tsx"),("Access boundary", "Server-validated demonstration session, protected root and simulation API, logout", "lib/demo-auth.ts; app/api/login/route.ts; app/api/logout/route.ts"),("Delivery", "Locked source manifest and compact Vercel deployment bundle", "smartcross-2.2.lock.json; scripts/build-deploy-bundle.mjs")])
para(doc, "The application is intentionally local-first. No live device adapter, external database, or outbound responder integration is present in the as-built prototype.")

doc.add_heading("3 Runtime sequence", 1)
make_table(doc, ["Step", "State or action", "Result"], [("1", "Authenticate at /access", "Server validates the controlled demonstration session before protected routes are served."),("2", "Select mode and scenario", "Normal or School Crossing selects a deterministic scenario contract."),("3", "Run condition", "Motion frames, timing, signal state, pedestrian state, and event records are created."),("4", "Observe seven stages", "Approach, pedestrian demand, sensor confirmation, traffic stop, protected WALK, route clear, traffic resumes."),("5", "Review evidence", "CCTV/radar labels, synthetic plate evidence, violation history, and incident controls remain simulation records."),("6", "Generate report", "The same in-memory session records are formatted into a readable PDF.")])

doc.add_heading("4 Crossing scene design", 1)
para(doc, "The scene uses a two-way mid-block road with upper and lower travel lanes, a horizontal road-spanning zebra crossing, tactile waiting areas, traffic signal heads at the side of the road, pedestrian foreground graphics, and a school courtyard only in School Crossing mode. Vehicles remain stationary at their initial approach markers until a run begins.")
make_table(doc, ["Element", "Normal Crossing", "School Crossing"], [("Location", "MPAJ mid-block local road", "School forecourt and direct crossing route"),("Pedestrian", "Single high-contrast pedestrian", "Supervised child group"),("Context", "Public footpath and streetscape", "School frontage, gate, and courtyard"),("Vehicle start", "Vehicle A lower/eastbound; Vehicle B upper/westbound", "Same lane assignment with school context"),("Detection", "Mirrored 40 m eastbound and westbound zones", "Same 40 m geometry with stricter school safeguards")])

doc.add_heading("5 Detection and evidence design", 1)
para(doc, "The nominal detection range is 40 metres upstream of each approach. This is a simulation and design baseline, not a field-certified measurement. Speed measurement is represented as a separate zone and should use calibrated radar, LiDAR, or an approved measurement sensor synchronised with CCTV evidence before field use.")
make_table(doc, ["Evidence element", "As-built behaviour", "Boundary"], [("CCTV", "Four logical views: EB incoming, EB outgoing, WB incoming, WB outgoing", "Visual simulation only; no live stream"),("Radar", "Approach and opposing radar labels support direction and tracking explanation", "No physical detector data"),("Speed zones", "Mirrored 40 m eastbound and westbound segmented zones", "Not enforcement-grade"),("ANPR", "Synthetic plate capture for authorised review", "No live lookup or identity inference"),("Sensor health", "Operational, constrained, or degraded conditions influence fallback", "Health is deterministic mock input")])

doc.add_heading("6 State machines", 1)
doc.add_heading("6.1 Standard crossing sequence", 2)
para(doc, "The seven-stage rail is the common contract for normal, school, violation, fallback, and incident records. A violation changes the action, signal, and review state while preserving chronological evidence.")
make_table(doc, ["Stage", "Meaning", "Safety expectation"], [("1 Approach", "Actors enter the configured approach context", "No pedestrian WALK before conflicting traffic is stopped."),("2 Pedestrian demand", "Pedestrian request or presence is evaluated", "WAIT/HOLD is retained when the route is unsafe."),("3 Sensors confirm", "Radar, CCTV, stop-line, and crossing state are correlated", "Uncertain inputs favour conservative handling."),("4 Both directions stop", "Traffic is held at RED or safe STOP", "No conflicting vehicle movement is released."),("5 Protected WALK", "Pedestrian moves through the marked route", "Clearance is not abruptly terminated."),("6 Route clears", "Pedestrian leaves the crossing and evidence is retained", "Vehicle release waits for route clear."),("7 Traffic resumes", "Normal or controlled recovery occurs", "Record closes with outcome and review status.")])
doc.add_heading("6.2 Emergency priority sequence", 2)
para(doc, "Emergency priority is higher than heavy-traffic optimization but does not bypass pedestrian clearance or signal safety. Verified ambulance, police, or fire-rescue requests are represented by this isolated state machine:")
para(doc, "NORMAL -> EMERGENCY_DETECTED -> VALIDATING -> DETERMINE_APPROACH -> CALCULATE_OCCUPANCY -> PEDESTRIAN_CLEARANCE -> STOP_NEW_ENTRY -> TRAFFIC_CLEARANCE -> EMERGENCY_PRIORITY -> EMERGENCY_PASSING -> CONFIRM_CROSSING_CLEAR -> RECOVERY -> NORMAL")
make_table(doc, ["State", "Decision and visible evidence"], [("EMERGENCY_DETECTED", "CCTV + radar candidate enters the 40 m approach range."),("VALIDATING", "Vehicle class, direction, source, confidence, and validation status are recorded."),("DETERMINE_APPROACH", "Eastbound travels left-to-right in the lower lane; Westbound travels right-to-left in the upper lane."),("CALCULATE_OCCUPANCY", "Pedestrian and crossing occupancy are assessed before priority."),("PEDESTRIAN_CLEARANCE", "People already crossing retain WALK and sufficient clearance."),("STOP_NEW_ENTRY", "New pedestrian entry is held while the corridor is prepared."),("TRAFFIC_CLEARANCE", "Conflicting vehicles move to their own stop line and hold STOP."),("EMERGENCY_PRIORITY / PASSING", "The verified responder receives a simulated green corridor in its approach lane."),("CONFIRM_CROSSING_CLEAR", "Outgoing coverage and crossing state support exit confirmation."),("RECOVERY", "Deferred demand and normal timing are restored gradually.")])

doc.add_heading("7 Priority and safety rules", 1)
for item in ["Collision risk and life-safety hazards outrank all other actions.","Verified emergency priority outranks adaptive heavy-traffic behaviour.","A pedestrian already inside the crossing retains protected clearance.","Unvalidated emergency-like detections remain in safe hold and cannot automatically create priority.","Vehicle signals are directional in the UI, but the existing two-way mid-block design does not invent conflicting physical movements.","A pedestrian WALK state is never granted while conflicting vehicles remain moving through the crossing.","Critical incidents require human review and do not dispatch police, MPAJ, fire, or hospital services.","Camera degradation invokes the conservative fallback instead of treating missing data as clear road space."]:
    bullet(doc, item)

doc.add_heading("8 Data contracts and records", 1)
make_table(doc, ["Record", "Key fields"], [("Scenario", "id, label, summary, mode, road user, severity, violation category, detection, verification, response, outcome"),("Motion frame", "timeline stage, vehicle A/B x and lane, heading, speed, state, pedestrian progress/state, signal states, controller, effect, route clear"),("Emergency priority", "vehicle type, direction, CCTV + radar source, confidence, validation status, detection distance, speed, priority status, frames, event log"),("Event log", "session ID, event ID, timestamp, category, scenario label, message, review state"),("CCTV evidence", "logical camera ID, direction, plate placeholder, detection reason, confidence, no-live-lookup boundary"),("PDF report", "session metadata, scenario outcome, timeline, event records, safety boundary, generated timestamp")])

doc.add_heading("9 Access, privacy, and deployment", 1)
para(doc, "The application is protected by a server-side demonstration login and signed session. Credentials are intentionally not repeated in this document. The production project is Vercel-hosted and published through the stable SmartCross alias. The repository lock and deployment bundle prevent accidental drift and exclude unrelated Phase 2 implementation files.")
para(doc, "The system stores no live faces, registration identities, or enforcement records. Synthetic evidence demonstrates review workflows. Any future operational deployment would require a separate identity design, data-retention policy, security review, lawful authority, and approved device integration.")

doc.add_heading("10 Verification and as-built acceptance", 1)
make_table(doc, ["Verification area", "As-built result"], [("Automated unit tests", "62 tests passing across simulation, access, text, and emergency-priority contracts."),("TypeScript", "No type errors in the source workspace."),("Lock verification", "33-file SmartCross baseline verified before bundle generation."),("Production build", "Vercel production builds reach READY status."),("Direction checks", "Eastbound left-to-right lower lane; Westbound right-to-left upper lane; opposing traffic yields at its own stop line."),("Regression boundary", "Existing violation catalogue, seven-stage timeline, fallback, PDF reports, login gate, and simulation-only boundary retained.")])

doc.add_heading("11 Limitations and commissioning requirements", 1)
para(doc, "SmartCross is not field-ready traffic infrastructure. Before real deployment, MPAJ and the relevant road authority would need to validate road geometry, signal phasing, camera lens and mounting calculations, radar or LiDAR calibration, lighting and weather performance, communications, controller interlocks, cybersecurity, privacy and retention, accessibility, emergency-service coordination, maintenance, fail-safe operation, and Malaysian regulatory compliance.")
para(doc, "The 40 m range, speed values, confidence values, timing values, and emergency progression are configurable simulation assumptions. They must not be presented as certified thresholds without site survey, engineering approval, and supervised commissioning evidence.")

doc.add_heading("12 Source traceability", 1)
make_table(doc, ["Design topic", "Source of implementation"], [("Scenario and event contracts", "lib/simulation.ts; lib/types.ts"),("Motion and lane direction", "lib/motion.ts; components/crossing-view.tsx"),("Emergency priority", "lib/emergency-priority.ts; components/dashboard.tsx"),("CCTV and detection styling", "components/crossing-view.tsx; app/enhancements.css"),("Fallback and degraded operation", "lib/camera-fallback.ts; components/dashboard.tsx"),("PDF reports", "lib/pdf-report.ts"),("Authentication", "lib/demo-auth.ts; app/access/page.tsx; app/api/login/route.ts"),("Deployment control", "smartcross-2.2.lock.json; scripts/verify-smartcross-2-2-lock.mjs; scripts/build-deploy-bundle.mjs")])

doc.add_heading("Appendix A Bloom validation mapping", 1)
make_table(doc, ["Level", "SmartCross evidence"], [("Remember", "Identify CCTV, radar, 40 m zones, signals, pedestrian states, emergency states, and safety boundaries."),("Understand", "Explain why normal timing, degraded fallback, heavy traffic, and emergency priority differ."),("Apply", "Run a selected Normal or School Crossing scenario and observe the seven stages."),("Analyze", "Compare event history, direction, signal state, sensor health, and pedestrian protection."),("Evaluate", "Review whether the response met STOP-before-WALK, human review, and recovery rules."),("Create", "Use traceable records to propose validated site-specific profiles for future commissioning.")])

doc.save(OUT)
print(OUT.resolve())
